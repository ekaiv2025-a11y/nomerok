import openpyxl, json, re, sys, collections, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from place_addr import addr_ru, guess_tags, ka_plain
SRC=sys.argv[1]; OUT=sys.argv[2]
wb=openpyxl.load_workbook(SRC,read_only=True)
CAT={'Еда и заведения':'food','Проживание':'stay','Продукты и напитки':'groceries','Красота и уход':'beauty','Медицина и аптеки':'health','Образование':'education','Одежда и обувь':'clothes','Розница':'shops','Дом и строительство':'home','Авто и транспорт':'auto','Культура и развлечения':'culture','Профессиональные услуги':'business','Недвижимость':'realty','Финансы':'finance','Спорт и фитнес':'sport','Электроника и связь':'electronics','Туризм и экскурсии':'tourism','События и услуги':'events','Животные':'pets','Цветы и подарки':'flowers','Мастерские и ремонт':'repair','Доставка и логистика':'delivery','Детские товары':'kids','Книги и хобби':'hobby','Коворкинги и офисы':'coworking','Медиа и творчество':'media','Производство и мастерские':'production','Аренда и прокат':'rent','Бытовые услуги':'services'}
TAG={'place_of_learning':'Учебное заведение','dance_club':'Танцевальный клуб','education':'Обучение','music_venue':'Концертная площадка','casino':'Казино','lounge':'Лаунж-бар','arts_and_entertainment':'Досуг и развлечения','preschool':'Детский сад','specialty_school':'Специализированная школа','college_university':'Колледж, университет','sport_league':'Спортивный клуб','health_care':'Медицинские услуги','family_service':'Семейные услуги','sports_and_recreation':'Спорт и отдых','fitness_centre':'Фитнес-центр','sports_centre':'Спортивный центр','indoor_play':'Детская игровая','ticket':'Билетная касса','outpost':'Пункт выдачи заказов','e-cigarette':'Электронные сигареты','mall':'Торговый центр','jewelry':'Ювелирный','fashion_accessories':'Аксессуары','museum':'Музей','gallery':'Галерея','electronics_repair':'Ремонт электроники','hvac':'Кондиционеры и вентиляция','company':'Офис компании','theatre':'Театр','cinema':'Кинотеатр','arts_centre':'Арт-центр','swimming_pool':'Бассейн','playground':'Детская площадка','pitch':'Спортивная площадка','stadium':'Стадион','shoemaker':'Ремонт обуви','tailor':'Ателье','carpenter':'Столярная мастерская','photographer':'Фотограф','watchmaker':'Ремонт часов','locksmith':'Ключи и замки','key_cutter':'Изготовление ключей','beauty':'Салон красоты','hairdresser':'Парикмахерская','car_repair':'Автосервис','car_wash':'Автомойка','tyres':'Шиномонтаж','bicycle':'Велосипеды','mobile_phone':'Мобильные телефоны','computer':'Компьютеры','copyshop':'Копицентр','laundry':'Прачечная','dry_cleaning':'Химчистка','travel_agency':'Турагентство','estate_agent':'Агентство недвижимости','lawyer':'Юрист','notary':'Нотариус','accountant':'Бухгалтер','insurance':'Страхование','it':'IT-компания','government':'Госучреждение','ngo':'Некоммерческая организация','school':'Школа','kindergarten':'Детский сад','library':'Библиотека','language_school':'Языковая школа','music_school':'Музыкальная школа','driving_school':'Автошкола','dancing_school':'Школа танцев','gym':'Тренажёрный зал','yoga':'Йога','spa':'Спа','massage':'Массаж','nightclub':'Ночной клуб','bowling_alley':'Боулинг','escape_game':'Квест','amusement_arcade':'Игровые автоматы','water_park':'Аквапарк','zoo':'Зоопарк','aquarium':'Аквариум','attraction':'Достопримечательность','viewpoint':'Смотровая площадка','marina':'Причал','beach_resort':'Пляж'}
def sub(s):
    if not s: return None
    s=str(s).strip()
    m=re.match(r'^(.+?):\s*([a-z_\-]+)$',s)
    if m: return TAG.get(m.group(2)) or m.group(1)
    if re.match(r'^[a-z_\-]+$',s): return TAG.get(s)
    return s
KA=dict(zip('აბგდევზთიკლმნოპჟრსტუფქღყშჩცძწჭხჯჰ',['a','b','g','d','e','v','z','t','i','k','l','m','n','o','p','zh','r','s','t','u','p','k','gh','q','sh','ch','ts','dz','ts','ch','kh','j','h']))
RU=dict(zip('абвгдеёжзийклмнопрстуфхцчшщъыьэюя',['a','b','v','g','d','e','e','zh','z','i','y','k','l','m','n','o','p','r','s','t','u','f','h','ts','ch','sh','sch','','y','','e','yu','ya']))
def translit(s):
    s=str(s).lower(); o=''
    for ch in s: o+=KA.get(ch, RU.get(ch, ch))
    o=re.sub(r'[^a-z0-9]+','-',o).strip('-')
    return o[:50].strip('-')
def nm(s):
    s=re.sub(r'^[^\w\d]+','',str(s).strip())
    s=re.sub(r'^["«“]+|["»”]+$','',s).strip()
    s=s.replace('""','"')
    return s or None
def clean(v):
    if v is None: return None
    v=str(v).strip()
    return v or None
def split(v):
    v=clean(v)
    if not v: return []
    return [x.strip() for x in re.split(r'[;\n|,]\s*(?=\+|https?://|\d)|[;\n|]',v) if x.strip()]
def phones(v):
    out=[]
    for p in split(v):
        d=re.sub(r'[^\d+]','',p)
        if len(re.sub(r'\D','',d))>=9 and d not in out: out.append(d if d.startswith('+') else ('+995'+d[-9:] if len(d)==9 else '+'+d))
    return out
def sites(v):
    return [u for u in split(v) if re.match(r'https?://',u) and not re.search(r'g\.co/kgs|google\.|goo\.gl',u)]
out=[]; slugs=set()
REPD=json.load(open(sys.argv[3])) if len(sys.argv)>3 else {}
def addslug(base,id_):
    s=(base or 'place')+'-'+id_[-4:].lower()
    s=re.sub(r'-+','-',s)
    i=2
    while s in slugs: s=f'{s}-{i}'; i+=1
    slugs.add(s); return s
ws=wb['Организации']; it=ws.iter_rows(values_only=True); h=next(it); ix={k:i for i,k in enumerate(h)}
for r in it:
    if not r[0]: continue
    g=lambda k: clean(r[ix[k]])
    src=g('Название источника'); ru=g('Название RU'); en=g('Название EN'); ka=g('Название KA')
    name=nm(ru or en or src)
    if not name: continue
    if re.search('[\u10A0-\u10FF]',name) and not re.search('[A-Za-zА-Яа-я]',name):
        ka=ka or name; name=ka_plain(name)
    en=nm(en) if en else en; ka=nm(ka) if ka else ka
    ig=[u for u in split(g('Instagram')) if 'instagram' in u]
    ph=phones(g('Телефоны источников')) or phones(g('Телефон с сайта'))
    st=sites(g('Сайты')); so=[u for u in split(g('Другие соцсети')) if re.match(r'https?://',u)]
    lat=g('Широта'); lng=g('Долгота')
    p={'id':r[0],'slug':addslug(translit(en or src or name) or translit(name), r[0]),'name':name,'name_en':en if en and en!=name else None,'name_ka':ka if ka and ka!=name else None,
       'cat':CAT.get(g('Категория RU'),'services'),'sub':sub(g('Подкатегория источника')),'addr':g('Адрес RU') or g('Адрес источника'),
       'lat':float(lat) if lat else None,'lng':float(lng) if lng else None,'phones':ph,'site':st[0] if st else None,'ig':ig[0] if ig else None,'socials':so[:3],
       'email':g('Деловой email'),'hours':g('Часы источника'),'desc_ru':g('Описание RU'),'desc_en':g('Описание EN'),
       'src':g('Источники') or 'OSM','area':'outskirts' if 'Окрестности' in (g('Территория') or '') else 'batumi'}
    p['addr_ru']=addr_ru(p.get('addr'))
    if p['cat'] in ('food','groceries'): p['tags']=guess_tags(' '.join(x for x in [src,ru,en,ka,g('Подкатегория источника')] if x), p.get('sub'))
    p['status']='published' if (ph or st or ig or so) else 'hidden'
    out.append({k:v for k,v in p.items() if v not in (None,[],'')})
# черновики ремонта телефонов (проверены по открытым источникам)
ws=wb['Описания RU EN']; it=ws.iter_rows(values_only=True); h=next(it); ix={k:i for i,k in enumerate(h)}
for r in it:
    if not r or not r[0]: continue
    g=lambda k: clean(r[ix[k]])
    name=g('Название'); contact=g('Контакт') or ''
    addr=re.sub(r'\s+—.*$','',g('Адрес источника') or '')
    p={'id':r[0],'slug':addslug(translit(name),r[0]),'name':name,'cat':'repair','sub':'Ремонт телефонов','addr':('Батуми, '+addr) if addr else None,
       'addr_ru':('Батуми, '+addr) if addr else None,'tags':['ремонт телефонов'],'phones':phones(contact),'desc_ru':REPD.get(r[0]) or g('Описание RU'),'desc_en':g('Description EN'),'src':'Сверка публичных источников','status':'published',
       'whatsapp':bool(re.search('whatsapp',contact,re.I)),'checked':str(r[ix['Дата сверки']])[:10] if r[ix['Дата сверки']] else None,'links':[u for u in split(g('Доказательства')) if re.match(r'https?://',u)][:4]}
    out.append({k:v for k,v in p.items() if v not in (None,[],'',False)})
json.dump(out,open(OUT,'w'),ensure_ascii=False,separators=(',',':'))
subs=collections.defaultdict(collections.Counter)
for p in out:
    if p.get('sub') and p['status']=='published': subs[p['cat']][p['sub']]+=1
json.dump({c:[s for s,n in v.most_common(12) if n>=3] for c,v in subs.items()},open(os.path.join(os.path.dirname(OUT),'place-subs.json'),'w'),ensure_ascii=False,indent=0)
c=collections.Counter(p['status'] for p in out); print(len(out),c)
print(collections.Counter(p.get('sub') for p in out).most_common(15))
